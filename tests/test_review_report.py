"""Review gate, finding library and render: edge cases found in review (no Salesforce org, no AI).

Run: python -m unittest discover -s tests -v      (from the repo root)
"""

import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import unittest
import zipfile
from pathlib import Path

from openpyxl import load_workbook

REPO = Path(__file__).resolve().parent.parent
SCRIPTS = REPO / "scripts"
FIXTURE = REPO / "tests" / "fixtures" / "seeded-org"
sys.path.insert(0, str(SCRIPTS))
sys.path.insert(0, str(REPO / "tests"))

from common import Run, load_config  # noqa: E402

OBJ_FILE = "metadata/force-app/main/default/objects/Invoice__c/Invoice__c.object-meta.xml"
CLS = "metadata/force-app/main/default/classes/InvoiceService.cls"


def script(name, *args):
    """Run a script; return (exit code, stderr)."""
    proc = subprocess.run([sys.executable, str(SCRIPTS / name), *map(str, args)], cwd=SCRIPTS,
                          capture_output=True, text=True, encoding="utf-8")
    return proc.returncode, proc.stderr


def finding(n, **kw):
    f = {"id": f"F-{n:03d}", "severity": "High", "area": "apex", "category": "performance",
         "title": f"InvoiceService runs SOQL inside a loop ({n})", "component": "ApexClass:InvoiceService",
         "related_components": [], "evidence": [{"path": CLS, "line_start": 1, "line_end": 1, "quote": "x"}],
         "impact": "Hits governor limits", "recommendation": "Bulkify the query", "effort": "S",
         "basis": "llm-observation", "fingerprint": f"fp{n}"}
    f.update(kw)
    return f


class ReviewReportTest(unittest.TestCase):
    def setUp(self):
        if not FIXTURE.exists():
            script("../tests/make_fixture.py")
        self.tmp = Path(tempfile.mkdtemp(prefix="sfhealth-rr-"))
        self.lib = self.tmp / "lib" / "library.jsonl"

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    # ------------------------------------------------------------------ helpers
    def make_run(self, client="Acme Widgets Inc", findings=None, date="2026-01-01") -> Path:
        run = self.tmp / re.sub(r"[^A-Za-z0-9_-]+", "-", client) / date
        shutil.copytree(FIXTURE, run)
        for sub in ("findings", "shards", "review", "report"):
            shutil.rmtree(run / sub, ignore_errors=True)
        obj = run / OBJ_FILE
        obj.write_text(obj.read_text(encoding="utf-8").replace("<label>Invoice</label>", "<label>Zephyr Ledger</label>"),
                       encoding="utf-8")
        code, err = script("build_index.py", "--run", run)
        self.assertEqual(code, 0, err)
        r = Run(run)
        r.set_meta(client=client, dataPolicy="metadata-only")
        r.log_step("preflight", "ok", 0, org={"username": "jdoe.harrington@acmewidgets.com",
                                              "instanceUrl": "https://acmewidgets.my.salesforce.com"})
        (run / "client-config.yml").write_text(
            f"dataPolicy: metadata-only\nlibrary:\n  path: '{self.lib.as_posix()}'\n", encoding="utf-8")
        (run / "findings").mkdir(exist_ok=True)
        (run / "findings" / "findings.json").write_text(json.dumps(
            {"findings": findings or [finding(1), finding(2), finding(3)]}), encoding="utf-8")
        return run

    def sheet(self, run):
        path = run / "review" / "findings_review.xlsx"
        wb = load_workbook(path)
        ws = wb["Findings"]
        cols = {c.value: c.column for c in ws[1]}
        return path, wb, ws, cols

    def export_and_decide(self, run, decisions, edit=None):
        code, err = script("review_gate.py", "export", "--run", run)
        self.assertEqual(code, 0, err)
        path, wb, ws, cols = self.sheet(run)
        for r in range(2, ws.max_row + 1):
            fid = ws.cell(r, cols["ID"]).value
            if fid in decisions:
                ws.cell(r, cols["Decision"]).value = decisions[fid]
            for (eid, col), val in (edit or {}).items():
                if eid == fid:
                    ws.cell(r, cols[col]).value = val
        wb.save(path)
        return path

    def accepted(self, run):
        return json.loads((run / "findings" / "accepted.json").read_text(encoding="utf-8"))

    # ------------------------------------------------------------------ review gate import
    def test_duplicate_id_rows_exit_2(self):
        run = self.make_run()
        path = self.export_and_decide(run, {"F-001": "Accept", "F-002": "Accept", "F-003": "Accept"})
        _, wb, ws, _ = self.sheet(run)
        ws.append([c.value for c in ws[2]])
        wb.save(path)
        code, err = script("review_gate.py", "import", "--run", run)
        self.assertEqual(code, 2)
        self.assertIn("F-001", err)
        self.assertIn("more than one row", err)

    def test_deleted_row_counts_as_pending(self):
        run = self.make_run()
        path = self.export_and_decide(run, {"F-001": "Accept", "F-002": "Accept", "F-003": "Accept"})
        _, wb, ws, cols = self.sheet(run)
        row = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, cols["ID"]).value == "F-002")
        ws.delete_rows(row)
        wb.save(path)
        code, err = script("review_gate.py", "import", "--run", run)
        self.assertEqual(code, 3, err)
        self.assertIn("no row", err)
        code, err = script("review_gate.py", "import", "--run", run, "--allow-pending")
        self.assertEqual(code, 0, err)
        acc = self.accepted(run)
        self.assertEqual(acc["review"]["pending"], 1)
        self.assertNotIn("F-002", {f["id"] for f in acc["findings"]})

    def test_accept_with_edits_is_applied_as_edit_and_reject_edits_ignored(self):
        run = self.make_run()
        self.export_and_decide(run, {"F-001": "Accept", "F-002": "reject ", "F-003": "Accept"},
                               {("F-001", "New severity"): "low", ("F-001", "Edited title"): "Better title",
                                ("F-002", "Edited title"): "ignored"})
        code, err = script("review_gate.py", "import", "--run", run)
        self.assertEqual(code, 0, err)
        self.assertIn("imported as Edit", err)
        self.assertIn("were ignored", err)
        f1 = next(f for f in self.accepted(run)["findings"] if f["id"] == "F-001")
        self.assertEqual((f1["severity"], f1["original_severity"], f1["title"]), ("Low", "High", "Better title"))
        self.assertEqual(f1["review"]["status"], "edited")
        self.assertEqual([r["id"] for r in self.accepted(run)["rejected"]], ["F-002"])

    def test_wrong_active_sheet_still_imports(self):
        run = self.make_run()
        path = self.export_and_decide(run, {"F-001": "Accept", "F-002": "Accept", "F-003": "Accept"})
        wb = load_workbook(path)
        wb.active = wb.sheetnames.index("How to review")
        wb.save(path)
        code, err = script("review_gate.py", "import", "--run", run)
        self.assertEqual(code, 0, err)
        self.assertEqual(len(self.accepted(run)["findings"]), 3)

    def test_numeric_and_blank_cells_and_unknown_decision(self):
        run = self.make_run()
        self.export_and_decide(run, {"F-001": "Edit", "F-002": "   ", "F-003": "Accept"},
                               {("F-001", "Edited title"): 12345, ("F-001", "Edited recommendation"): 3.5})
        code, err = script("review_gate.py", "import", "--run", run)
        self.assertEqual(code, 3, err)                       # whitespace decision = pending
        code, err = script("review_gate.py", "import", "--run", run, "--allow-pending")
        self.assertEqual(code, 0, err)
        f1 = next(f for f in self.accepted(run)["findings"] if f["id"] == "F-001")
        self.assertEqual((f1["title"], f1["recommendation"]), ("12345", "3.5"))
        self.export_and_decide(run, {"F-001": "Maybe", "F-002": 7, "F-003": "Accept"})
        code, err = script("review_gate.py", "import", "--run", run)
        self.assertEqual(code, 2)
        self.assertIn("unknown decision 'Maybe'", err)
        self.assertIn("unknown decision '7'", err)

    # ------------------------------------------------------------------ formula injection
    def assert_no_formulas(self, xlsx: Path):
        with zipfile.ZipFile(xlsx) as z:
            for name in z.namelist():
                if name.startswith("xl/worksheets/sheet"):
                    self.assertNotIn("<f>", z.read(name).decode("utf-8"), name)

    def test_formula_injection_is_written_as_text(self):
        evil = '=HYPERLINK("http://evil.example/x","click")'
        run = self.make_run(findings=[finding(1, title=evil, recommendation="+cmd|' /C calc'!A0"),
                                      finding(2, impact="@SUM(1+1)", component="-2+3")])
        code, err = script("review_gate.py", "export", "--run", run)
        self.assertEqual(code, 0, err)
        _, _, ws, cols = self.sheet(run)
        cell = ws.cell(2, cols["Title"])
        self.assertEqual((cell.value, cell.data_type), (evil, "s"))
        self.assert_no_formulas(run / "review" / "findings_review.xlsx")

        code, err = script("review_gate.py", "accept-all", "--run", run)
        self.assertEqual(code, 0, err)
        code, err = script("render.py", "--run", run, "--formats", "xlsx")
        self.assertEqual(code, 0, err)
        out = run / "report" / "baseline.xlsx"
        self.assert_no_formulas(out)
        ws = load_workbook(out)["Findings"]
        self.assertEqual(ws.cell(2, 4).value, evil)
        self.assertEqual(ws.cell(2, 4).data_type, "s")

    # ------------------------------------------------------------------ finding library
    def test_library_scrubs_client_data_and_hides_notes_across_clients(self):
        import library
        run = self.make_run()
        cfg = load_config()
        cfg["library"]["path"] = str(self.lib)
        f = {"title": 'AWI: Zephyr Ledger "Project Falcon" calls https://api.acmewidgets.com/v1 and '
                      "acmewidgets.my.salesforce.com, owned by Harrington, AcmeWidgets and Widgets team",
             "area": "apex", "category": "security", "severity": "High",
             "recommendation": "Ask jdoe.harrington@acmewidgets.com", "basis": "llm-observation"}
        note = "Harrington at Acme says the Zephyr Ledger callout is retired"
        library.record(Run(run), cfg, [(f, "rejected", note, "Already fixed"), (f, "rejected", note, "Already fixed")])
        saved = self.lib.read_text(encoding="utf-8")
        for secret in ("AWI", "Zephyr", "Ledger", "Project Falcon", "acmewidgets", "AcmeWidgets", "Widgets",
                       "Harrington", "jdoe", "Acme", "https://"):
            self.assertNotIn(secret, saved, secret)
        self.assertNotIn(note, saved)
        entry = json.loads(saved.splitlines()[0])
        self.assertEqual(entry["reasonCategory"], "Already fixed")
        self.assertRegex(entry["client"], r"^[0-9a-f]{12}$")

        same = library.Hints(cfg, Run(run))(f)
        self.assertIn("This client's last reject note", same)
        other = self.make_run(client="Globex", findings=[finding(1)])
        hint = library.Hints(cfg, Run(other))({**f, "title": f["title"]})
        self.assertTrue(hint.startswith("Seen 2x before"), hint)
        self.assertIn("already fixed", hint)                  # category only
        self.assertNotIn("note", hint)
        self.assertNotIn("retired", hint)
        self.assertNotIn("retired", "".join(library.false_positives(cfg, "apex")))

    def test_library_disabled_in_run_config(self):
        run = self.make_run()
        (run / "client-config.yml").write_text(
            f"dataPolicy: metadata-only\nlibrary:\n  enabled: false\n  path: '{self.lib.as_posix()}'\n", encoding="utf-8")
        self.export_and_decide(run, {"F-001": "Reject", "F-002": "Accept", "F-003": "Accept"})
        code, err = script("review_gate.py", "import", "--run", run)
        self.assertEqual(code, 0, err)
        self.assertFalse(self.lib.exists())

    # ------------------------------------------------------------------ render
    def accept_and_synthesize(self, run, cited):
        code, err = script("review_gate.py", "accept-all", "--run", run)
        self.assertEqual(code, 0, err)
        syn = {"kind": "sf-health-synthesis",
               "executive_summary": {"headline": "h", "paragraphs": ["a", "b"],
                                     "top_recommendations": [{"text": "t", "finding_ids": cited}]},
               "area_summaries": [], "roadmap": {"quick_wins": [{"title": "Fix", "finding_ids": cited,
                                                                 "effort": "S", "why": "w"}],
                                                 "short_term": [], "long_term": []}}
        p = run / "findings" / "sf-health-synthesis.json"
        p.write_text(json.dumps(syn), encoding="utf-8")
        later = time.time() + 5
        os.utime(p, (later, later))
        return p

    def test_render_escapes_embedded_json(self):
        nasty = "Bad <!--<script> and </script><b>x</b> &   end"
        run = self.make_run(findings=[finding(1, title=nasty, impact="<!--"), finding(2), finding(3)])
        self.accept_and_synthesize(run, ["F-001", "F-002", "F-003"])
        code, err = script("render.py", "--run", run, "--formats", "html")
        self.assertEqual(code, 0, err)
        page = (run / "report" / "report.html").read_text(encoding="utf-8")
        m = re.search(r'<script id="report-data" type="application/json">(.*?)</script>', page, re.S)
        blob = m.group(1)
        for bad in ("<", ">", "&", " "):
            self.assertNotIn(bad, blob)
        data = json.loads(blob)
        self.assertEqual(data["findings"][0]["title"] if data["findings"][0]["id"] == "F-001"
                         else next(f["title"] for f in data["findings"] if f["id"] == "F-001"), nasty)

    def test_render_refuses_stale_synthesis(self):
        run = self.make_run()
        self.accept_and_synthesize(run, ["F-001", "F-999"])
        code, err = script("render.py", "--run", run, "--formats", "docx")
        self.assertEqual(code, 2)
        self.assertIn("F-999", err)
        code, err = script("render.py", "--run", run, "--formats", "docx", "--allow-stale-synthesis")
        self.assertEqual(code, 0, err)
        from docx import Document
        text = "\n".join(p.text for p in Document(run / "report" / "report.docx").paragraphs)
        self.assertNotIn("F-999", text)
        self.assertIn("F-001", text)

    def test_render_refuses_synthesis_older_than_review(self):
        run = self.make_run()
        p = self.accept_and_synthesize(run, ["F-001", "F-002", "F-003"])
        code, err = script("render.py", "--run", run, "--formats", "html")
        self.assertEqual(code, 0, err)
        old = time.time() - 3600
        os.utime(p, (old, old))                              # summary written before the latest decisions
        code, err = script("render.py", "--run", run, "--formats", "html")
        self.assertEqual(code, 2)
        self.assertIn("older than", err)


if __name__ == "__main__":
    unittest.main()
