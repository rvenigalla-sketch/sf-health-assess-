"""Deterministic pipeline tests on the synthetic seeded org (no Salesforce org, no AI).

Run: python -m unittest discover -s tests -v      (from the repo root)
"""

import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
SCRIPTS = REPO / "scripts"
FIXTURE = REPO / "tests" / "fixtures" / "seeded-org"
sys.path.insert(0, str(SCRIPTS))

import plan_shards  # noqa: E402
from common import load_config  # noqa: E402
from common import Run  # noqa: E402
from score import assessed_areas, score_areas  # noqa: E402


def run_script(name, *args):
    proc = subprocess.run([sys.executable, str(SCRIPTS / name), *map(str, args)], cwd=SCRIPTS,
                          capture_output=True, text=True, encoding="utf-8")
    if proc.returncode != 0:
        raise AssertionError(f"{name} failed ({proc.returncode}):\n{proc.stderr[-2000:]}")
    return proc


class PipelineTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if not FIXTURE.exists():
            run_script("../tests/make_fixture.py")
        cls.tmp = Path(tempfile.mkdtemp(prefix="sfhealth-"))
        cls.run_dir = cls.tmp / "Fixture" / "2026-01-01"
        shutil.copytree(FIXTURE, cls.run_dir)
        run_script("build_index.py", "--run", cls.run_dir)
        run_script("plan_shards.py", "--run", cls.run_dir)
        cls.inv = {i["id"]: i for i in json.loads((cls.run_dir / "index" / "inventory.json").read_text(encoding="utf-8"))}
        cls.auto = json.loads((cls.run_dir / "index" / "automation_map.json").read_text(encoding="utf-8"))

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    # ---------------- index
    def test_inventory_types(self):
        self.assertTrue(self.inv["ApexClass:InvoiceServiceTest"]["isTest"])
        self.assertFalse(self.inv["ApexClass:InvoiceService"]["isTest"])
        self.assertEqual(self.inv["ApexTrigger:InvoiceTrigger"]["object"], "Invoice__c")
        self.assertIn("InvoiceService", self.inv["ApexTrigger:InvoiceTrigger"]["refs"]["classes"])

    def test_automation_map_and_overlaps(self):
        a = self.auto["Invoice__c"]
        self.assertEqual(len(a["triggers"]), 1)
        self.assertEqual(len(a["flowsAfterSave"]), 2)
        self.assertEqual(len(a["processBuilders"]), 1)
        self.assertEqual(len(a["workflowRules"]), 1)
        writers = a["fieldOverlaps"]
        self.assertIn("Amount_Total__c", writers)
        self.assertTrue(any(w.startswith("WorkflowRule:") for w in writers["Amount_Total__c"]))
        self.assertIn("Status__c", writers)
        self.assertTrue(any(w.startswith("ProcessBuilder:") for w in writers["Status__c"]))

    def test_flow_digest_flags(self):
        digest = (self.run_dir / self.inv["Flow:Invoice_Line_Sync"]["digestPath"]).read_text(encoding="utf-8")
        self.assertIn("IN-LOOP:Each_Sibling", digest)
        self.assertIn("NO-FAULT-PATH", digest)
        self.assertIn("possible hard-coded record IDs", digest)

    def test_field_usage(self):
        table = (self.run_dir / "index" / "datamodel" / "Invoice__c.txt").read_text(encoding="utf-8")
        rows = {l.split(" | ")[0]: l.split(" | ") for l in table.splitlines() if " | " in l}
        self.assertEqual(rows["Legacy_Code__c"][-1], "0")
        self.assertNotEqual(rows["AmountTotal__c"][-1], "0")
        self.assertNotEqual(rows["Amount_Total__c"][-1], "0")

    def test_automation_summary_order(self):
        text = (self.run_dir / "index" / "automation" / "Invoice__c.txt").read_text(encoding="utf-8")
        self.assertLess(text.index("1. Before-save"), text.index("5. Apex AFTER triggers"))
        self.assertLess(text.index("7. Process Builders"), text.index("8. After-save"))
        self.assertIn("Amount_Total__c", text.split("FIELDS WRITTEN BY MORE THAN ONE")[1])

    # ---------------- shards
    def test_shards_cover_everything(self):
        shards = [json.loads(p.read_text(encoding="utf-8")) for p in (self.run_dir / "shards").glob("*.json")
                  if p.name != "plan.json"]
        covered = {c for s in shards for c in s["components"]}
        for cid in ("ApexClass:InvoiceService", "ApexClass:InvoiceSearchController", "ApexClass:InvoiceServiceTest",
                    "Flow:Invoice_Line_Sync", "FieldUsage:Invoice__c", "AutomationSummary:Invoice__c"):
            self.assertIn(cid, covered)
        for s in shards:
            self.assertLessEqual(s["sourceTokens"], 45000)

    def test_pack_limits_and_balance(self):
        items = [{"tokens": t} for t in (30, 30, 30, 30, 5, 5, 5)]
        bins = plan_shards.pack([[i] for i in items], budget=70, cost=lambda i: i["tokens"], max_items=3)
        for b in bins:
            self.assertLessEqual(sum(i["tokens"] for i in b), 70)
            self.assertLessEqual(len(b), 3)
        self.assertEqual(sum(len(b) for b in bins), len(items))
        self.assertEqual(len(bins), 3)

    # ---------------- validator
    def test_validator_rejects_fabrication(self):
        raw = self.run_dir / "findings" / "raw"
        raw.mkdir(parents=True, exist_ok=True)
        path = self.inv["ApexClass:InvoiceService"]["path"]
        base = {"area": "apex", "category": "bulkification", "severity": "High", "component": "ApexClass:InvoiceService",
                "evidence_ids": [], "basis": "llm-observation", "confidence": "high", "impact": "x",
                "recommendation": "y", "effort": "S", "related_components": []}
        good = {**base, "title": "real", "evidence": [{"path": path, "line_start": 4, "line_end": 4,
                "quote": "List<Invoice__c> siblings = [SELECT Id, Amount__c FROM Invoice__c"}]}
        bad_quote = {**base, "title": "bad quote", "evidence": [{"path": path, "line_start": 4, "line_end": 4,
                     "quote": "Database.delete(everything);"}]}
        bad_file = {**base, "title": "bad file", "evidence": [{"path": "metadata/nope.cls", "line_start": 1,
                    "line_end": 1, "quote": "public class"}]}
        bad_id = {**base, "title": "bad id", "basis": "tool-confirmed", "evidence_ids": ["CA-99999"],
                  "evidence": good["evidence"]}
        escape = {**base, "title": "escape", "evidence": [{"path": "../../../../Windows/win.ini", "line_start": 1,
                  "line_end": 1, "quote": "fonts"}]}
        shard_id = next(p.stem for p in (self.run_dir / "shards").glob("apex-0*.json"))
        (raw / f"{shard_id}.json").write_text(json.dumps({"shard_id": shard_id,
            "findings": [good, bad_quote, bad_file, bad_id, escape],
            "coverage": {"components_reviewed": ["ApexClass:InvoiceService"], "skipped": []}, "follow_up": []}),
            encoding="utf-8")
        run_script("validate_findings.py", "--run", self.run_dir)
        ok = [json.loads(l) for l in (self.run_dir / "findings" / "validated.jsonl").read_text(encoding="utf-8").splitlines()]
        bad = [json.loads(l) for l in (self.run_dir / "findings" / "rejected.jsonl").read_text(encoding="utf-8").splitlines()]
        self.assertEqual([f["title"] for f in ok], ["real"])
        self.assertEqual(sorted(f["title"] for f in bad), ["bad file", "bad id", "bad quote", "escape"])
        cov = json.loads((self.run_dir / "findings" / "coverage.json").read_text(encoding="utf-8"))
        self.assertTrue(cov["shards"][shard_id]["missing"], "unreviewed components must be reported")

        # Score, review gate (with real decisions) and render on top of the validated finding.
        run_script("score.py", "--run", self.run_dir)
        run_script("review_gate.py", "export", "--run", self.run_dir)
        from openpyxl import load_workbook
        xlsx = self.run_dir / "review" / "findings_review.xlsx"
        wb = load_workbook(xlsx)
        ws = wb.active
        from review_gate import EDITABLE_FROM
        ws.cell(2, EDITABLE_FROM).value = "Edit"
        ws.cell(2, EDITABLE_FROM + 1).value = "Medium"
        wb.save(xlsx)
        # Library off: tests must never write to the user's real finding library.
        no_lib = self.tmp / "no-library.yml"
        no_lib.write_text("library:\n  enabled: false\n", encoding="utf-8")
        run_script("review_gate.py", "import", "--run", self.run_dir, "--config", no_lib)
        accepted = json.loads((self.run_dir / "findings" / "accepted.json").read_text(encoding="utf-8"))
        self.assertEqual(accepted["findings"][0]["severity"], "Medium")
        self.assertEqual(accepted["findings"][0]["original_severity"], "High")

        accepted["findings"][0]["title"] = "</script><script>window.pwned=1</script>"
        (self.run_dir / "findings" / "accepted.json").write_text(json.dumps(accepted), encoding="utf-8")
        run_script("render.py", "--run", self.run_dir, "--client", "Fixture Co")
        html = (self.run_dir / "report" / "report.html").read_text(encoding="utf-8")
        self.assertNotIn("</script><script>window.pwned", html)
        for name in ("report.docx", "baseline.xlsx", "report.artifact.html"):
            self.assertTrue((self.run_dir / "report" / name).exists())

    # ---------------- scoring
    def test_areas_without_shards_not_assessed(self):
        areas = assessed_areas(Run(self.run_dir))
        self.assertIn("apex", areas)
        self.assertNotIn("crma", areas)          # fixture has no CRM Analytics assets
        self.assertNotIn("integration", areas)   # and no integration metadata

    def test_scoring_formula_and_critical_cap(self):
        cfg = load_config()
        f = lambda sev, area="apex": {"severity": sev, "area": area}
        s = score_areas([f("Low")], cfg, {"apex"})
        self.assertEqual(s["areas"]["apex"]["score"], 98)
        self.assertEqual(s["areas"]["flow"]["rating"], "Not assessed")
        s = score_areas([f("Critical")], cfg, {"apex", "flow", "security", "automation"})
        self.assertNotEqual(s["overallRating"], "Good")


CLASSES = Path("metadata/force-app/main/default/classes")


def jsonl(path: Path) -> list[dict]:
    return [json.loads(l) for l in path.read_text(encoding="utf-8").splitlines() if l.strip()] if path.exists() else []


class Phase2Test(unittest.TestCase):
    """Rule checks, incremental runs and the finding library (each test works on its own copy)."""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="sfhealth2-"))

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def snapshot(self, date: str) -> Path:
        run = self.tmp / "Fixture" / date
        shutil.copytree(FIXTURE, run)
        for sub in ("findings", "shards", "review", "report"):      # outputs from earlier pilots, if any
            shutil.rmtree(run / sub, ignore_errors=True)
        run_script("build_index.py", "--run", run)
        return run

    # ---------------- rule checks
    def test_rule_checks_cover_planted_problems(self):
        sys.path.insert(0, str(REPO / "tests"))
        from eval import matches
        run = self.snapshot("2026-01-01")
        run_script("rule_checks.py", "--run", run)
        run_script("validate_findings.py", "--run", run)
        found = jsonl(run / "findings" / "validated.jsonl")
        self.assertEqual(len(found), 3)
        self.assertEqual(jsonl(run / "findings" / "rejected.jsonl"), [])
        planted = {p["id"]: p for p in json.loads((REPO / "tests" / "fixtures" / "expected_findings.json")
                                                  .read_text(encoding="utf-8"))["planted"]}
        for pid in ("P13", "P14", "P15"):
            self.assertTrue(any(matches(f, planted[pid]) for f in found), pid)
        # Reviewers whose shard covers those components are told not to repeat them.
        run_script("plan_shards.py", "--run", run)
        flow_shards = [json.loads(p.read_text(encoding="utf-8")) for p in (run / "shards").glob("flow-*.json")]
        self.assertTrue(any(s["context"].get("alreadyReported") for s in flow_shards))

    # ---------------- incremental
    def test_incremental_reviews_only_changes_and_carries_the_rest(self):
        prev = self.snapshot("2026-01-01")
        search = (prev / CLASSES / "InvoiceSearchController.cls").read_text(encoding="utf-8").splitlines()
        line = next(n for n, l in enumerate(search, 1) if "without sharing" in l)
        ev = lambda path, n, q: [{"path": path, "line_start": n, "line_end": n, "quote": q}]
        base = {"area": "apex", "category": "security", "severity": "High", "evidence_ids": [], "basis": "llm-observation",
                "confidence": "high", "impact": "x", "recommendation": "y", "effort": "S", "related_components": []}
        (prev / "findings").mkdir(exist_ok=True)
        (prev / "findings" / "validated.jsonl").write_text("\n".join(json.dumps(f) for f in [
            {**base, "finding_id": "apex-01#1", "shard_id": "apex-01", "title": "Search runs without sharing",
             "component": "ApexClass:InvoiceSearchController",
             "evidence": ev((CLASSES / "InvoiceSearchController.cls").as_posix(), line, search[line - 1].strip())},
            {**base, "finding_id": "apex-01#2", "shard_id": "apex-01", "title": "SOQL in loop",
             "component": "ApexClass:InvoiceService",
             "evidence": ev((CLASSES / "InvoiceService.cls").as_posix(), 1, "public class InvoiceService")},
            {**base, "finding_id": "apex-02#1", "shard_id": "apex-02", "title": "Formatter quote no longer there",
             "component": "ApexClass:CleanFormatter",
             "evidence": ev((CLASSES / "CleanFormatter.cls").as_posix(), 1, "text that is not in the file")},
        ]) + "\n", encoding="utf-8")
        (prev / "findings" / "verdicts").mkdir()
        (prev / "findings" / "verdicts" / "apex-01__1.json").write_text(json.dumps(
            {"finding_id": "apex-01#1", "verdict": "confirmed", "corrected_severity": "High", "reason": "ok"}), encoding="utf-8")

        cur = self.snapshot("2026-02-01")
        svc = cur / CLASSES / "InvoiceService.cls"
        svc.write_text(svc.read_text(encoding="utf-8") + "\n// changed\n", encoding="utf-8")
        run_script("build_index.py", "--run", cur)
        run_script("plan_shards.py", "--run", cur, "--baseline", "auto")

        plan = json.loads((cur / "shards" / "plan.json").read_text(encoding="utf-8"))
        self.assertEqual(plan["incremental"]["baseline"], "2026-01-01")
        self.assertLess(plan["shards"], plan["incremental"]["fullPlanShards"])
        reviewed = {c for p in (cur / "shards").glob("*-*.json")
                    for c in json.loads(p.read_text(encoding="utf-8")).get("components", [])}
        self.assertIn("ApexClass:InvoiceService", reviewed)          # changed
        self.assertIn("ApexTrigger:InvoiceTrigger", reviewed)        # calls the changed class
        self.assertNotIn("ApexClass:InvoiceSearchController", reviewed)
        carried = json.loads((cur / "findings" / "raw" / "carried.json").read_text(encoding="utf-8"))["findings"]
        self.assertEqual([f["component"] for f in carried],
                         ["ApexClass:InvoiceSearchController", "ApexClass:CleanFormatter"])
        self.assertTrue((cur / "findings" / "verdicts" / "carried__1.json").exists())
        self.assertTrue(assessed_areas(Run(cur)) >= {"apex", "flow"})   # unchanged areas still count

        run_script("validate_findings.py", "--run", cur)
        ok = {f["finding_id"] for f in jsonl(cur / "findings" / "validated.jsonl")}
        bad = {f["finding_id"] for f in jsonl(cur / "findings" / "rejected.jsonl")}
        self.assertIn("carried#1", ok)
        self.assertIn("carried#2", bad)                              # stale quote is caught, not carried blindly
        args = json.loads(run_script("workflow_args.py", "verify", "--run", cur).stdout)
        self.assertNotIn("carried#1", {f["finding_id"] for f in args["findings"]})   # keeps last run's verdict

    # ---------------- finding library
    def test_library_strips_names_and_gives_hints(self):
        import library
        run = self.snapshot("2026-01-01")
        cfg = load_config()
        cfg["library"]["path"] = str(self.tmp / "lib.jsonl")
        f = {"title": "InvoiceService runs SOQL on Invoice__c inside a loop (admin@fixture.com, Fixture)",
             "area": "apex", "category": "performance", "severity": "High",
             "recommendation": "Move the query in InvoiceService out of the loop", "basis": "llm-observation"}
        library.record(Run(run), cfg, [(f, "rejected", "Fixture says InvoiceService is retired"),
                                       (f, "rejected", "same")])
        saved = (self.tmp / "lib.jsonl").read_text(encoding="utf-8")
        for secret in ("InvoiceService", "Invoice__c", "admin@fixture.com", "Fixture"):
            self.assertNotIn(secret, saved)
        self.assertIn("<name>", saved)
        hint = library.Hints(cfg, Run(run))({**f, "title": "InvoiceSearchController runs SOQL on Invoice__c inside a loop"})
        self.assertTrue(hint.startswith("Seen 2x before"), hint)
        self.assertEqual(len(library.false_positives(cfg, "apex")), 1)


if __name__ == "__main__":
    unittest.main()
