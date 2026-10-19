"""Stage 6 - Consultant review gate. Nothing reaches the client report without a human decision.

export: findings/findings.json -> review/findings_review.xlsx
        One row per finding, with Decision (Accept / Reject / Edit) and optional edits.
import: review/findings_review.xlsx -> findings/accepted.json
        Keeps Accept and Edit rows (with edits applied), recomputes scores from accepted findings only.
        Fails (exit 3) if any finding is still undecided or its row was deleted, unless --allow-pending
        (pending findings are left out). Fails (exit 2) on duplicate ID rows, unknown decisions or
        invalid severities.

The client config saved in the run folder is reused when --config is omitted.

Usage:
  python review_gate.py export --run <run> [--config client.yml]
  python review_gate.py import --run <run> [--allow-pending] [--config client.yml]
  python review_gate.py accept-all --run <run>     # testing only: accept everything as-is
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
import time
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

import library
from common import Run, earlier_run, info, load_run_config, read_json, write_json
from score import SEV_ORDER, assessed_areas, score_areas

SHEET = "Findings"
COLUMNS = [
    ("ID", 8), ("Severity", 10), ("Area", 12), ("Title", 50), ("Component", 34), ("Evidence", 46),
    ("Impact", 50), ("Recommendation", 50), ("Effort", 7), ("Verification", 26), ("Baseline", 12),
    ("Library hint", 34),
    ("Decision", 11), ("New severity", 12), ("Edited title", 40), ("Edited recommendation", 50),
    ("Reject reason", 22), ("Reviewer notes", 40),
]
EDITABLE_FROM = 13          # 1-based column index of "Decision"
EDIT_COLUMNS = ("New severity", "Edited title", "Edited recommendation")
DECISIONS = {"accept": "Accept", "accepted": "Accept", "reject": "Reject", "rejected": "Reject",
             "edit": "Edit", "edited": "Edit"}
SEV_FILL = {"Critical": "F8D7DA", "High": "FDE2C8", "Medium": "FFF3CD", "Low": "E2E8F0"}
# Leading characters a spreadsheet app may treat as the start of a formula (CSV/DDE injection).
FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


# --------------------------------------------------------------------------- spreadsheet safety
def xl_safe(cell) -> None:
    """Force a cell holding org / AI text to be plain text, never a formula. openpyxl turns any string
    starting with '=' into a formula; the quote prefix also stops Excel re-parsing it when the cell is edited."""
    v = cell.value
    if isinstance(v, str) and v.startswith(FORMULA_START):
        cell.data_type = "s"
        try:
            cell.quotePrefix = True
        except (AttributeError, TypeError):     # older openpyxl: data_type alone already stores text
            pass


def xl_append(ws, row: list) -> None:
    """ws.append for rows containing org / AI data: every cell goes through xl_safe."""
    ws.append(row)
    for cell in ws[ws.max_row]:
        xl_safe(cell)


def cell_text(v) -> str:
    """Spreadsheet cell -> stripped text (numbers and dates included; None -> '')."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def accepted_hash(findings: list[dict]) -> str:
    """Identifies the accepted set (ids and the fields the summary uses); stored in accepted.json."""
    key = sorted([f["id"], f["severity"], f["area"], f["component"], f["title"], f.get("recommendation", "")]
                 for f in findings)
    return hashlib.sha256(json.dumps(key, ensure_ascii=False).encode("utf-8")).hexdigest()[:16]


def write_accepted(run: Run, payload: dict) -> None:
    """Write accepted.json with the accepted-set hash. acceptedChangedAt only moves when the set changes,
    so render.py can tell whether the summary was written before the latest decisions."""
    path = run.path("findings", "accepted.json")
    old = (read_json(path, {}) or {}).get("review", {})
    h = accepted_hash(payload["findings"])
    payload["review"]["acceptedHash"] = h
    payload["review"]["acceptedChangedAt"] = (old.get("acceptedChangedAt") if old.get("acceptedHash") == h
                                              and old.get("acceptedChangedAt") else time.time())
    write_json(path, payload)


# --------------------------------------------------------------------------- export
def previous_decisions(run: Run) -> dict[str, str]:
    """fingerprint -> what the consultant decided in the previous reviewed run of this client."""
    prev = earlier_run(run, "findings/accepted.json")
    if not prev:
        return {}
    acc = read_json(prev / "findings" / "accepted.json", {})
    if (acc.get("review") or {}).get("autoAccepted"):
        return {}
    out = {f["fingerprint"]: f.get("review", {}).get("status", "accepted") for f in acc.get("findings", []) if f.get("fingerprint")}
    out.update({f["fingerprint"]: "rejected" for f in acc.get("rejected", []) if f.get("fingerprint")})
    return {fp: f"{prev.name}: {status}" for fp, status in out.items()}


def hint_text(f: dict, prev: dict, hints) -> str:
    parts = []
    if f.get("fingerprint") in prev:
        parts.append(f"Last run {prev[f['fingerprint']]}")
    if f.get("carried_from"):
        parts.append("Unchanged since last run (not re-reviewed)")
    lib = hints(f) if hints else ""
    if lib:
        parts.append(lib)
    return ". ".join(parts)[:400]


def export(run: Run, cfg: dict) -> Path:
    data = read_json(run.path("findings", "findings.json"))
    prev = previous_decisions(run)
    hints = library.Hints(cfg, run)
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append([c for c, _ in COLUMNS])
    for i, (_, w) in enumerate(COLUMNS, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        ws.cell(1, i).font = Font(bold=True)
        if i >= EDITABLE_FROM:
            ws.cell(1, i).fill = PatternFill("solid", fgColor="DBEAFE")
    for f in data["findings"]:
        ev = "\n".join(f"{e['path'].split('/')[-1]}:{e['line_start']}-{e['line_end']}" for e in f["evidence"])
        ver = f.get("verification", {})
        xl_append(ws, [f["id"], f["severity"], f["area"], f["title"], f["component"], ev, f["impact"], f["recommendation"],
                       f["effort"], f"{ver.get('verdict', '-')}: {ver.get('reason', '')}"[:300] if ver else "-",
                       f.get("baseline_status", ""), hint_text(f, prev, hints)] + [""] * (len(COLUMNS) - 12))
        r = ws.max_row
        ws.cell(r, 2).fill = PatternFill("solid", fgColor=SEV_FILL[f["severity"]])
        for c in range(1, len(COLUMNS) + 1):
            ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")
    last = max(ws.max_row, 2)
    names = [c for c, _ in COLUMNS]
    letter = lambda name: ws.cell(1, names.index(name) + 1).column_letter
    for name, choices in (("Decision", ["Accept", "Reject", "Edit"]), ("New severity", list(SEV_ORDER)),
                          ("Reject reason", library.REJECT_REASONS)):
        dv = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{letter(name)}2:{letter(name)}{last}")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(COLUMNS)).column_letter}{ws.max_row}"

    help_ws = wb.create_sheet("How to review")
    for line in [
        "For every row, set Decision:",
        "  Accept - correct as written.",
        "  Reject - wrong or not relevant: pick a Reject reason, and say why in Reviewer notes.",
        "  Edit   - correct, but change severity / title / recommendation in the blue columns.",
        "  (An Accept row with anything in the edit columns is imported as Edit; edits on a Reject row are ignored.)",
        "Don't delete rows: a finding without a row counts as undecided.",
        "Check the Evidence column against the code before accepting High and Critical findings.",
        "Library hint: how similar findings were decided in past assessments (all clients, names removed),",
        "  and what you decided for this exact finding last time. A hint, not a decision: still check the row.",
        "  Reviewer notes are kept (names removed) only as hints for this same client; other clients see",
        "  just the Reject reason category.",
        "Verification 'rule check' rows are computed by a script from the metadata, not by the AI.",
        "Save the file, then run: python review_gate.py import --run <run folder>",
    ]:
        help_ws.append([line])
    help_ws.column_dimensions["A"].width = 100

    # For audit only: findings the verifier disproved. They are not imported; if one is actually
    # right, raise it with the team and re-run the verifier rather than editing files by hand.
    refuted_path = run.path("findings", "refuted.jsonl")
    refuted = [json.loads(l) for l in refuted_path.open(encoding="utf-8") if l.strip()] if refuted_path.exists() else []
    if refuted:
        rws = wb.create_sheet("Refuted by verifier")
        rws.append(["Severity", "Component", "Title", "Verifier reason"])
        for i, w in enumerate((10, 34, 60, 100), 1):
            rws.column_dimensions[rws.cell(1, i).column_letter].width = w
            rws.cell(1, i).font = Font(bold=True)
        for f in refuted:
            xl_append(rws, [f["severity"], f["component"], f["title"], (f.get("verification") or {}).get("reason", "")])
            for c in range(1, 5):
                rws.cell(rws.max_row, c).alignment = Alignment(wrap_text=True, vertical="top")
    out = run.path("review", "findings_review.xlsx")
    wb.save(out)
    return out


# --------------------------------------------------------------------------- import
def findings_sheet(wb):
    """The 'Findings' sheet, whichever sheet was active when the file was saved. Falls back to the first
    sheet whose header has ID and Decision (in case the tab was renamed)."""
    if SHEET in wb.sheetnames:
        return wb[SHEET]
    for ws in wb.worksheets:
        if {"ID", "Decision"} <= {cell_text(c.value) for c in ws[1]}:
            return ws
    return None


def do_import(run: Run, cfg: dict, allow_pending: bool) -> int:
    data = read_json(run.path("findings", "findings.json"))
    by_id = {f["id"]: f for f in data["findings"]}
    ws = findings_sheet(load_workbook(run.path("review", "findings_review.xlsx")))
    if ws is None:
        info(f"REVIEW: no '{SHEET}' sheet (or sheet with ID and Decision columns) in review/findings_review.xlsx")
        return 2
    header = [cell_text(c.value) for c in ws[1]]
    col = {name: header.index(name) for name, _ in COLUMNS if name in header}
    missing_cols = [c for c in ("ID", "Decision") if c not in col]
    if missing_cols:
        info(f"REVIEW: the '{ws.title}' sheet has no {', '.join(missing_cols)} column")
        return 2
    get = lambda row, name: cell_text(row[col[name]]) if name in col and col[name] < len(row) else ""

    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if get(r, "ID")]
    counts = Counter(get(r, "ID") for r in rows)
    dupes = sorted(i for i, n in counts.items() if n > 1)
    if dupes:
        info(f"REVIEW: these IDs appear on more than one row: {', '.join(dupes)}. "
             "Keep one row per finding (delete the copies), then import again.")
        return 2

    accepted, rejected, pending, problems, notes_out = [], [], [], [], []
    seen = set()
    for row in rows:
        fid = get(row, "ID")
        if fid not in by_id:
            notes_out.append(f"row with unknown ID {fid!r} ignored")
            continue
        seen.add(fid)
        f = dict(by_id[fid])
        raw = get(row, "Decision")
        if not raw:
            pending.append(fid)
            continue
        decision = DECISIONS.get(raw.lower())
        if not decision:
            problems.append(f"{fid}: unknown decision {raw!r} (use Accept, Reject or Edit)")
            continue
        notes, reason = get(row, "Reviewer notes"), get(row, "Reject reason")
        edits = {c: get(row, c) for c in EDIT_COLUMNS if get(row, c)}
        if decision == "Reject":
            if edits:
                notes_out.append(f"{fid}: rejected, so the edits in {', '.join(edits)} were ignored")
            f["review"] = {"status": "rejected", "notes": notes, "reason": reason}
            rejected.append(f)
            continue
        if decision == "Accept" and edits:
            notes_out.append(f"{fid}: marked Accept but has edits in {', '.join(edits)}; imported as Edit")
            decision = "Edit"
        f["review"] = {"status": "accepted" if decision == "Accept" else "edited", "notes": notes}
        if decision == "Edit":
            new_sev = edits.get("New severity", "").capitalize()
            if new_sev:
                if new_sev not in SEV_ORDER:
                    problems.append(f"{fid}: unknown severity {new_sev!r}")
                    continue
                f["original_severity"] = f.get("original_severity", f["severity"])
                f["severity"] = new_sev
            if edits.get("Edited title"):
                f["title"] = edits["Edited title"]
            if edits.get("Edited recommendation"):
                f["recommendation"] = edits["Edited recommendation"]
        accepted.append(f)
    no_row = [i for i in by_id if i not in seen]
    pending += no_row

    for n in notes_out:
        info(f"REVIEW NOTE: {n}")
    if problems:
        for p in problems:
            info(f"REVIEW: {p}")
        return 2
    if pending and not allow_pending:
        info(f"REVIEW: {len(pending)} finding(s) still have no decision (e.g. {', '.join(pending[:5])})"
             + (f"; {len(no_row)} of them have no row in the sheet (deleted?)" if no_row else "")
             + ". Finish the review or pass --allow-pending.")
        return 3
    accepted.sort(key=lambda f: (SEV_ORDER[f["severity"]], f["area"], f["component"]))
    write_accepted(run, {
        "findings": accepted, "scores": score_areas(accepted, cfg, assessed_areas(run)),
        "review": {"accepted": len(accepted), "rejected": len(rejected), "pending": len(pending)},
        "baseline": data.get("baseline"),
        # Kept so the next run's sheet can show "last run: rejected" (not used in the report).
        "rejected": [{k: f.get(k) for k in ("id", "fingerprint", "title", "component", "severity", "review")} for f in rejected],
    })
    saved = library.record(run, cfg, [(f, f["review"]["status"], f["review"].get("notes", ""), f["review"].get("reason", ""))
                                      for f in accepted + rejected])
    run.log_step("review-import", "ok", 0, accepted=len(accepted), rejected=len(rejected), pending=len(pending),
                 librarySaved=saved)
    info(f"REVIEW: {len(accepted)} accepted, {len(rejected)} rejected, {len(pending)} pending -> findings/accepted.json"
         + (f"; {saved} decision(s) saved to the finding library (names removed)" if saved else ""))
    return 0


def accept_all(run: Run, cfg: dict) -> int:
    data = read_json(run.path("findings", "findings.json"))
    for f in data["findings"]:
        f["review"] = {"status": "accepted", "notes": "auto-accepted (testing)"}
    write_accepted(run, {
        "findings": data["findings"], "scores": score_areas(data["findings"], cfg, assessed_areas(run)),
        "review": {"accepted": len(data["findings"]), "rejected": 0, "pending": 0, "autoAccepted": True},
        "baseline": data.get("baseline"),
    })
    info("REVIEW: all findings auto-accepted (testing only; the report will say so)")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("action", choices=["export", "import", "accept-all"])
    ap.add_argument("--run", required=True)
    ap.add_argument("--allow-pending", action="store_true")
    ap.add_argument("--config", help="client config (default: the copy saved in the run folder)")
    args = ap.parse_args()
    run = Run(Path(args.run))
    cfg = load_run_config(run, args.config)
    if args.action == "export":
        info(f"REVIEW: wrote {export(run, cfg)}")
        return 0
    if args.action == "accept-all":
        return accept_all(run, cfg)
    return do_import(run, cfg, args.allow_pending)


if __name__ == "__main__":
    sys.exit(main())
