"""Stage 7 - Render the client deliverables. Deterministic.

Inputs : findings/accepted.json (after the review gate), findings/sf-health-synthesis.json,
         index/metrics.json, index/automation_map.json, findings/baseline_diff.json,
         index/datamodel/*.txt and index/orgfacts/*.txt (baseline workbook)
Outputs: report/report.html            interactive dashboard (open locally, or attach)
         report/report.artifact.html   same page without the document wrapper, for publishing as an artifact
         report/report.docx            client report (structure of the A360 sample report)
         report/baseline.xlsx          baseline workbook: findings, data model, security and access, user access, automation

The summary (findings/sf-health-synthesis.json) must match the current accepted findings: render stops
(exit 2) if it cites a finding that isn't accepted, or was written before the review decisions last
changed. Re-run the summary, or pass --allow-stale-synthesis to render anyway (citations of findings
that aren't accepted are then dropped and reported as warnings).

The client config saved in the run folder is reused (score thresholds).

Usage:
  python render.py --run <run> [--client "Display Name"] [--formats html,docx,xlsx] [--allow-stale-synthesis]
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

from common import PLUGIN_ROOT, Run, Timer, info, load_run_config, read_json
from review_gate import xl_append
from diff_baseline import METRIC_PATHS, dig

SEVS = ["Critical", "High", "Medium", "Low"]


# --------------------------------------------------------------------------- shared model
class StaleSynthesis(Exception):
    pass


def check_synthesis(run: Run, synthesis: dict, accepted: dict) -> list[str]:
    """Reasons the summary doesn't belong to the current accepted findings (empty = fine)."""
    if not synthesis:
        return []
    ids = {f["id"] for f in accepted["findings"]}
    road = synthesis.get("roadmap", {})
    cited = [i for k in ("quick_wins", "short_term", "long_term") for it in road.get(k, []) for i in it.get("finding_ids", [])]
    cited += [i for r in (synthesis.get("executive_summary") or {}).get("top_recommendations", [])
              for i in r.get("finding_ids", [])]
    out = []
    unknown = sorted(set(cited) - ids)
    if unknown:
        out.append(f"summary cites finding ids that are not accepted: {unknown[:10]}")
    rev = accepted.get("review", {})
    if synthesis.get("acceptedHash") and rev.get("acceptedHash") and synthesis["acceptedHash"] != rev["acceptedHash"]:
        out.append("summary was written for a different set of accepted findings")
    changed = rev.get("acceptedChangedAt")
    syn_path = run.path("findings", "sf-health-synthesis.json")
    if changed and syn_path.exists() and syn_path.stat().st_mtime < float(changed):
        out.append("summary is older than the latest change to the accepted findings (review_gate.py import ran after it)")
    return out


def keep_accepted(synthesis: dict, ids: set) -> None:
    """Drop citations of findings that are not accepted (roadmap items left with none are removed)."""
    road = synthesis.get("roadmap", {})
    for k in ("quick_wins", "short_term", "long_term"):
        if k in road:
            for it in road[k]:
                it["finding_ids"] = [i for i in it.get("finding_ids", []) if i in ids]
            road[k] = [it for it in road[k] if it["finding_ids"]]
    for r in (synthesis.get("executive_summary") or {}).get("top_recommendations", []):
        r["finding_ids"] = [i for i in r.get("finding_ids", []) if i in ids]


def load_model(run: Run, client: str | None, cfg: dict | None = None, allow_stale: bool = False) -> dict:
    accepted = read_json(run.path("findings", "accepted.json"))
    if not accepted:
        raise SystemExit("findings/accepted.json missing: run review_gate.py import (or accept-all for a draft)")
    synthesis = read_json(run.path("findings", "sf-health-synthesis.json"), {})
    metrics = read_json(run.path("index", "metrics.json"), {})
    automation = read_json(run.path("index", "automation_map.json"), {})
    log = read_json(run.log_path, {})
    meta = log.get("meta", {})
    org = next((s.get("org") for s in log.get("steps", []) if s.get("step") == "preflight" and s.get("org")), {}) or {}
    diff = read_json(run.path("findings", "baseline_diff.json"), {}) or {}
    ids = {f["id"] for f in accepted["findings"]}
    cfg = cfg or load_run_config(run)

    # The summary may only describe the accepted findings it was written from.
    problems = check_synthesis(run, synthesis, accepted)
    if problems and not allow_stale:
        raise StaleSynthesis("; ".join(problems))
    keep_accepted(synthesis, ids)
    road = synthesis.get("roadmap", {})
    cited = [i for k in ("quick_wins", "short_term", "long_term") for it in road.get(k, []) for i in it["finding_ids"]]
    unplaced = sorted(ids - set(cited))
    if synthesis and unplaced:
        problems.append(f"{len(unplaced)} accepted finding(s) not placed in the roadmap: {unplaced[:10]}")
    if synthesis and unplaced:
        road.setdefault("long_term", []).append({"title": "Other accepted findings", "finding_ids": unplaced,
                                                 "effort": "M", "why": "Not grouped by the summary; schedule with the client."})

    counts = metrics.get("componentCounts", {})
    scope = ", ".join(f"{counts.get(k, 0)} {label}" for k, label in (
        ("ApexClass", "Apex classes"), ("ApexTrigger", "triggers"), ("Flow", "flows"), ("CustomObject", "objects"),
        ("CustomField", "fields"), ("Profile", "profiles"), ("PermissionSet", "permission sets"),
        ("WaveDashboard", "CRM Analytics dashboards"), ("WaveRecipe", "recipes")) if counts.get(k))
    facts = [{"label": label, "value": dig(metrics, p)} for label, p in METRIC_PATHS if dig(metrics, p) is not None]
    return {
        "meta": {"client": client or meta.get("client") or run.root.parent.name, "date": run.root.name,
                 "orgHost": (org.get("instanceUrl") or "").replace("https://", ""), "dataPolicy": meta.get("dataPolicy", "full"),
                 "scope": scope},
        "findings": accepted["findings"], "scores": accepted["scores"], "review": accepted.get("review", {}),
        "synthesis": synthesis, "automation": automation, "facts": facts,
        "thresholds": {k: cfg["scoring"][k] for k in ("good", "fair")},
        "baseline": diff if diff.get("previous") else None, "problems": problems,
    }


# --------------------------------------------------------------------------- HTML
def logo_svg() -> str:
    """A360 logo, inlined: XML prolog and comments removed, class names prefixed so they can't
    collide with the page's CSS, and labelled for screen readers."""
    p = PLUGIN_ROOT / "templates" / "a360-logo.svg"
    if not p.exists():
        return ""
    svg = p.read_text(encoding="utf-8")
    svg = re.sub(r"<\?xml.*?\?>|<!--.*?-->", "", svg, flags=re.S).strip()
    svg = re.sub(r"\.st(\d+)", r".a360-st\1", svg)
    svg = re.sub(r'class="st(\d+)"', r'class="a360-st\1"', svg)
    svg = re.sub(r"<svg\b[^>]*?viewBox=\"([^\"]+)\"[^>]*>",
                 r'<svg class="a360-logo" viewBox="\1" role="img" aria-label="A360" xmlns="http://www.w3.org/2000/svg">', svg, count=1)
    return svg


def script_json(obj) -> str:
    """JSON safe to embed in a <script> element: no <, > or & survives, so org text can't open a tag or
    a comment ('<!--<script' would blank the page) or end the script. U+2028/2029 are escaped too."""
    data = json.dumps(obj, ensure_ascii=False)
    for ch, code in (("<", "\\u003c"), (">", "\\u003e"), ("&", "\\u0026"),
                     ("\u2028", "\\u2028"), ("\u2029", "\\u2029")):
        data = data.replace(ch, code)
    return data


def render_html(model: dict, out_dir: Path) -> list[Path]:
    template = (PLUGIN_ROOT / "templates" / "report.html").read_text(encoding="utf-8")
    data = script_json({k: v for k, v in model.items() if k != "problems"})
    title = f"{model['meta']['client']} Org Health"
    page = (template.replace("__TITLE__", title.replace("<", "")).replace("__A360_LOGO__", logo_svg())
            .replace("__REPORT_DATA__", data))
    artifact = out_dir / "report.artifact.html"
    artifact.write_text(page, encoding="utf-8")
    full = out_dir / "report.html"
    full.write_text('<!doctype html>\n<html lang="en">\n<head>\n<meta charset="utf-8">\n'
                    '<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">\n'
                    "</head>\n<body>\n" + page + "\n</body>\n</html>\n", encoding="utf-8")
    return [full, artifact]


# --------------------------------------------------------------------------- Word
def render_docx(model: dict, out_dir: Path) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    SEV_RGB = {"Critical": "B42318", "High": "C2410C", "Medium": "946200", "Low": "52606D"}
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10.5)

    def shade(cell, hex_fill):
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_fill)
        tc_pr.append(shd)

    def table(headers, rows, widths=None):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        for i, h in enumerate(headers):
            c = t.rows[0].cells[i]
            c.text = h
            c.paragraphs[0].runs[0].bold = True
            shade(c, "E9EDF1")
        for r in rows:
            cells = t.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = "" if v is None else str(v)
        return t

    m, sc, syn = model["meta"], model["scores"], model["synthesis"]
    accepted_ids = {f["id"] for f in model["findings"]}
    ex = syn.get("executive_summary", {})
    doc.add_heading(f"{m['client']} – Salesforce Org Health Assessment", level=0)
    p = doc.add_paragraph(f"Assessment date: {m['date']}")
    if model["review"].get("autoAccepted"):
        warn = doc.add_paragraph()
        run_ = warn.add_run("DRAFT: findings have not been reviewed by a consultant. Do not send to the client.")
        run_.bold = True
        run_.font.color.rgb = RGBColor.from_string("B42318")

    doc.add_heading("1. Executive summary", level=1)
    doc.add_paragraph(f"Overall org health: {sc['overallRating']} ({sc['overall'] if sc['overall'] is not None else '-'} / 100). "
                      + ex.get("headline", ""))
    for para in ex.get("paragraphs", []):
        doc.add_paragraph(para)
    counts = {s: sum(1 for f in model["findings"] if f["severity"] == s) for s in SEVS}
    table(["Severity", "Findings"], [[s, counts[s]] for s in SEVS])
    if ex.get("top_recommendations"):
        doc.add_paragraph().add_run("Top recommendations").bold = True
        for r in ex["top_recommendations"]:
            fids = [i for i in r.get("finding_ids", []) if i in accepted_ids]
            doc.add_paragraph(r["text"] + (f" ({', '.join(fids)})" if fids else ""), style="List Number")

    doc.add_heading("2. Scope", level=1)
    doc.add_paragraph("A360 reviewed a read-only copy of the org's metadata, retrieved with the Salesforce CLI"
                      + ("" if m["dataPolicy"] == "metadata-only" else ", plus read-only org checks (Security Health Check, sfdx-hardis)")
                      + ". Nothing in the org was changed. Every finding cites the exact lines it is based on and was "
                        "reviewed by an A360 consultant.")
    table(["Measure", "Value"], [[f["label"], f["value"]] for f in model["facts"]])

    doc.add_heading("3. Health scorecard", level=1)
    summaries = {a["area"]: a["summary"] for a in syn.get("area_summaries", [])}
    table(["Area", "Score", "Rating", "Summary"],
          [[a["label"], a["score"] if a["score"] is not None else "-", a["rating"], summaries.get(k, "")]
           for k, a in sc["areas"].items()])

    doc.add_heading("4. Findings register", level=1)
    doc.add_paragraph("Effort: S = under 1 day, M = 1–3 days, L = more than 3 days.")
    t = table(["ID", "Severity", "Area", "Finding", "Component", "Impact", "Recommendation", "Effort"],
              [[f["id"], f["severity"], sc["areas"].get(f["area"], {}).get("label", f["area"]), f["title"], f["component"],
                f["impact"], f["recommendation"], f["effort"]] for f in model["findings"]])
    for row in t.rows[1:]:
        sev = row.cells[1].text
        if sev in SEV_RGB:
            row.cells[1].paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(SEV_RGB[sev])
            row.cells[1].paragraphs[0].runs[0].bold = True

    doc.add_heading("5. Automation map", level=1)
    doc.add_paragraph("Everything that runs when a record is saved, per object (active only).")
    act = lambda xs: sum(1 for x in xs if x.get("status", "Active") == "Active")
    rows = []
    for obj, a in sorted(model["automation"].items(), key=lambda kv: -kv[1]["activeAutomationCount"]):
        rows.append([obj, sum(t_["active"] for t_ in a["triggers"]), act(a["flowsBeforeSave"]), act(a["flowsAfterSave"]),
                     act(a["processBuilders"]), sum(w["active"] for w in a["workflowRules"]),
                     sum(v["active"] for v in a["validationRules"]), ", ".join(a.get("fieldOverlaps", {}).keys()) or "-"])
    table(["Object", "Triggers", "Before-save flows", "After-save flows", "Process Builders", "Workflow rules",
           "Validation rules", "Fields written twice"], rows)

    doc.add_heading("6. Recommended roadmap", level=1)
    road = syn.get("roadmap", {})
    for key, label in (("quick_wins", "Quick wins (first 2–3 weeks)"), ("short_term", "Short term (1–2 months)"),
                       ("long_term", "Longer term (3–6 months)")):
        doc.add_heading(label, level=2)
        items = road.get(key, [])
        if not items:
            doc.add_paragraph("Nothing in this phase.")
        for it in items:
            fids = [i for i in it["finding_ids"] if i in accepted_ids]
            para = doc.add_paragraph(style="List Bullet")
            para.add_run(it["title"]).bold = True
            para.add_run(f" ({', '.join(fids) + '; ' if fids else ''}effort {it['effort']}). {it['why']}")

    doc.add_heading("7. Baseline and severity definitions", level=1)
    doc.add_paragraph("This report, the automation map and the metadata snapshot are the baseline for future "
                      "assessments. The baseline workbook (baseline.xlsx) lists the data model, security and access, "
                      "user access and automation as they stand today.")
    if model["baseline"]:
        b = model["baseline"]
        doc.add_paragraph(f"Compared with the assessment of {b['previous']}: {b['findings']['new']} new, "
                          f"{b['findings']['fixed']} fixed, {b['findings']['unchanged']} unchanged findings.")
    table(["Severity", "Definition"], [
        ["Critical", "Causing failures or exposing data today; fix immediately"],
        ["High", "Likely to cause errors, data issues or limit failures under normal use"],
        ["Medium", "Makes the org harder to maintain or less efficient"],
        ["Low", "Clean-up that reduces clutter and technical debt"]])
    doc.add_paragraph(f"Scoring: {sc['formula']['text']}.").alignment = WD_ALIGN_PARAGRAPH.LEFT
    out = out_dir / "report.docx"
    doc.save(out)
    return out


# --------------------------------------------------------------------------- Excel baseline
def _table_file(path: Path) -> tuple[list[str], list[list[str]]]:
    lines = [l for l in path.read_text(encoding="utf-8").splitlines() if l and not l.startswith("#")]
    if not lines:
        return [], []
    split = lambda l: [c.strip() for c in l.split(" | ")]
    return split(lines[0]), [split(l) for l in lines[1:]]


def render_xlsx(model: dict, run: Run, out_dir: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    wb.remove(wb.active)

    def sheet(name, headers, rows, widths=None):
        ws = wb.create_sheet(name[:31])
        xl_append(ws, headers)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="E9EDF1")
        for r in rows:
            xl_append(ws, r)
        for i, h in enumerate(headers, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = (widths or {}).get(h, min(60, max(12, len(h) + 4)))
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
        return ws

    ws = sheet("Findings", ["ID", "Severity", "Area", "Title", "Component", "Impact", "Recommendation", "Effort",
                            "Evidence", "Review", "Baseline status", "Fingerprint"],
               [[f["id"], f["severity"], f["area"], f["title"], f["component"], f["impact"], f["recommendation"], f["effort"],
                 "; ".join(f"{e['path']}:{e['line_start']}" for e in f["evidence"]), f.get("review", {}).get("status"),
                 f.get("baseline_status", ""), f.get("fingerprint")] for f in model["findings"]],
               {"Title": 50, "Impact": 50, "Recommendation": 50, "Evidence": 50, "Component": 34})
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    # Data model: every custom field with where-used counts.
    rows, headers = [], []
    for f in sorted(run.path("index", "datamodel").glob("*.txt")):
        h, rs = _table_file(f)
        headers = headers or ["Object", *h]
        rows += [[f.stem, *r] for r in rs]
    if rows:
        sheet("Data model", headers, rows, {"Field": 40, "UsedIn": 50})

    for topic, name in (("permissions", "Security and access"), ("sharing", "Sharing defaults"),
                        ("healthcheck", "Health Check"), ("privileged_users", "User access"),
                        ("licenses_users", "Users and licences"), ("connected_apps", "Connected apps")):
        p = run.path("index", "orgfacts", f"{topic}.txt")
        if p.exists():
            h, rs = _table_file(p)
            if h:
                sheet(name, h, rs, {"Tier0-1 permissions": 80, "Setting": 60})

    act = lambda xs: [x["name"] for x in xs if x.get("status", "Active") == "Active"]
    sheet("Automation map", ["Object", "Active count", "Before-save flows", "Apex triggers", "Validation rules",
                             "Workflow rules", "Process Builders", "After-save flows", "Fields written twice"],
          [[o, a["activeAutomationCount"], ", ".join(act(a["flowsBeforeSave"])),
            ", ".join(t["name"] for t in a["triggers"] if t["active"]),
            ", ".join(v["name"] for v in a["validationRules"] if v["active"]),
            ", ".join(w["name"] for w in a["workflowRules"] if w["active"]), ", ".join(act(a["processBuilders"])),
            ", ".join(act(a["flowsAfterSave"])), ", ".join(a.get("fieldOverlaps", {}).keys())]
           for o, a in sorted(model["automation"].items())])
    sheet("Metrics", ["Measure", "Value"], [[f["label"], f["value"]] for f in model["facts"]])
    out = out_dir / "baseline.xlsx"
    wb.save(out)
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--run", required=True)
    ap.add_argument("--client", help="client display name for the report (default: run folder's client)")
    ap.add_argument("--formats", default="html,docx,xlsx")
    ap.add_argument("--allow-stale-synthesis", action="store_true",
                    help="render even if the summary doesn't match the accepted findings (drops unknown ids)")
    ap.add_argument("--config", help="client config (default: the copy saved in the run folder)")
    args = ap.parse_args()
    run = Run(Path(args.run))
    formats = {f.strip() for f in args.formats.split(",")}
    cfg = load_run_config(run, args.config)

    with Timer() as t:
        try:
            model = load_model(run, args.client, cfg, args.allow_stale_synthesis)
        except StaleSynthesis as e:
            info(f"RENDER: stale summary: {e}. Re-run the summary (finalize step 2), "
                 "or pass --allow-stale-synthesis to render anyway.")
            run.log_step("render", "failed", t.seconds, error=str(e))
            return 2
        out_dir = run.path("report")
        written = []
        if "html" in formats:
            written += render_html(model, out_dir)
        if "docx" in formats:
            written.append(render_docx(model, out_dir))
        if "xlsx" in formats:
            written.append(render_xlsx(model, run, out_dir))

    for p in model["problems"]:
        info(f"RENDER WARNING: {p}")
    run.log_step("render", "ok", t.seconds, outputs=[p.name for p in written], warnings=model["problems"])
    info("RENDER: " + ", ".join(str(p) for p in written))
    return 0


if __name__ == "__main__":
    sys.exit(main())
